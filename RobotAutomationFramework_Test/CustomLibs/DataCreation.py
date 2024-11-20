import pandas as pd
import re
import xlsxwriter
import win32com.client as win32
import openpyxl
import subprocess

def create_data_excel(filepath):
    df1=pd.read_excel(filepath + "\\TestData.xlsx", "LoginPage", keep_default_na=False, dtype="object")
    df2=pd.read_excel(filepath + "\\TestData.xlsx", "Smoke", keep_default_na=False, dtype="object")
    #df2 = pd.read_excel(filepath + "\\TestData.xlsx", "Smoke", keep_default_na=False, dtype={"ColumnName":"string"})

    df1.set_index("TestCaseID", drop=True, inplace=True)
    df2.set_index("TestCaseID", drop=True, inplace=True)

    test1=df1.to_dict(orient="index")
    test2=df2.to_dict(orient="index")

    f = open(filepath + "\\TestData.robot", "w")
    print("Save to File")
    f.write("*** Variables ***\n")
    f.write("\n#######    LoginPage  ###########\n\n")
    f.write(str(test1))
    f.write("\n  ######## Smoke ################### \n\n")
    f.write(str(test2))
    f.close()

def format_data(filepath):
    fin = open(filepath + "\\TestData.robot", "rt")
    data = fin.read()
    data = re.sub("},", "\n", data)
    data = re.sub("}}", "\n", data)
    data = re.sub("{'", "", data)
    data = re.sub("':", "=", data)
    data = re.sub("}=", "} ", data)
    data = re.sub("= '", "=", data)
    data = re.sub("', '", "   ", data)
    data = re.sub(" '&", "&", data)
    data = re.sub("'\n", "\n", data)
    data = re.sub(", '", "  ", data)
    data = re.sub("= ", "=", data)
    data = re.sub('"', "", data)
    fin.close()
    fin = open(filepath + "\\TestData.robot", "wt")
    fin.write(data)
    fin.close()

def create_a_txt_file(filepath):
    file = open(filepath + "\\OutputData.txt", "w+")

def remove_whitespace(instring):
    return instring.strip

def write_to_file(filepath,header_key, header_values):
    file = open(filepath + "\\OutputData.txt", "r+")
    content = file.read()
    content_length = len(content)
    print(content_length)
    if content_length == 0:
        file.write(str(header_key) + "\n")
        file.write(str(header_values) + "\n")
    else:
        file.write(str(header_values) + "\n")
    file.close()

def format_output_data(filepath):
    file = open(filepath + "\\OutputData.txt", "r+")
    data = file.read()
    data = re.sub("[\[]", "",data)
    data = re.sub("]", "", data)
    data = re.sub("'", "", data)
    file.close()
    file = open(filepath + "\\OutputData.txt", "wt")
    file.write(data)
    file.close()

def write_to_excel(filepath, sheet_name, index):
    df = pd.read_table(filepath + "\\OutputData.txt", sep=",")
    df.set_index(index, drop=True, inplace=True)
    df.to_excel(filepath + "\\output.xlsx", sheet_name)

def format_excel(filepath,sheet_name):
    workbook = xlsxwriter.Workbook(filepath + "\\output.xlsx")
    worksheet = workbook.add_worksheet(sheet_name)
    worksheet.add_table('A1:F5',{'style' : 'Table Style Light 17'})

def  sort_Assc(List_sort):
    new_List = sorted(List_sort)
    print(type(new_List))
    return new_List

def string_convert(string_list):
    li = list(string_list.split("|"))
    print(type(li))
    return li

def  ConvertList_to_Upper(List_sort):
    for i in range(len(List_sort)):
        List_sort[i]= List_sort[i].upper()
    print(type(List_sort))
    return List_sort

def convert_xls_to_xlsx():
    fname = "\\\\LOUISILON02S\\USERDAT01\\AHS8367\\Downloads\\Reports.xls"
    excel = win32.gencache.EnsureDispatch('Excel.Application')
    wb = excel.Workbooks.Open(fname)
    wb.SaveAs(fname + "x", FileFormat=51)  # FileFormat = 51 is for .xlsx extension
    wb.Close()  # FileFormat = 56 is for .xls extension
    excel.Application.Quit()

def read_xlsfile(filepath):
    df = pd.read_excel('\\\\LOUISILON02S\\USERDAT01\\AHS8367\\Downloads\\Reports.xlsx', header=None)
    print(df)
    print(df.values[2][2])
    df26 = pd.read_excel(filepath + "\\TestData.xlsx", "TC12_Genetic", header=None)
    print(df26.values[2][11])

    global Total_Rows
    Total_Rows = df.shape[0]
    print(Total_Rows)

    for i in range(1, Total_Rows):
        if df.values[i][2] in df26.values[2][11]:
            # return True
            print(str(df.values[i][2])+"=> Found")
        else:
            return False
            # print(str(df.values[i][2])+" => Not Found")


def add_fax_Detail_into_TestData(filepath,Totalcount, Sheetname, ColNameArray, row_num, Parameter):
    wb = openpyxl.load_workbook(filepath + "\\TestData.xlsx")
    sheet = wb.get_sheet_by_name(Sheetname)
    for x in range(Totalcount):
        sheet[ColNameArray[x] + str(row_num)].value = Parameter[x]
    wb.save(filepath + "\\TestData.xlsx")

def cler_ie():
    commands = (
                "RunDll32.exe InetCpl.cpl,ClearMyTracksByProcess 255", # Deletes ALL History
                "RunDll32.exe InetCpl.cpl,ClearMyTracksByProcess 1",   # Deletes History Only
                "RunDll32.exe InetCpl.cpl,ClearMyTracksByProcess 2",   # Deletes Cookies Only
                "RunDll32.exe InetCpl.cpl,ClearMyTracksByProcess 8",   # Deletes Temporary Internet Files Only
                "RunDll32.exe InetCpl.cpl,ClearMyTracksByProcess 16",  # Deletes Form Data Only
                "RunDll32.exe InetCpl.cpl,ClearMyTracksByProcess 32",  # Deletes Password History Only
                )

    for command in commands:
        with subprocess.Popen(command) as p:
            p.wait()
            print(f"{p.returncode} - {command}")




