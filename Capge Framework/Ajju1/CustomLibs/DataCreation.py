import pandas as pd
import os
import sys
import fileinput
import re
import xlsxwriter
import xlrd
import pyexcel
import win32com.client as win32
import openpyxl
import subprocess

def create_data_excel():
    c = "C:\\DCOE_RobotAutomationFramework_Azure\\Data"
    df1 = pd.read_excel(c + "\\TestData.xlsx", "TS_Security", keep_default_na=False)
    df2 = pd.read_excel(c + "\\TestData.xlsx", "Smoke", keep_default_na=False)
    df3 = pd.read_excel(c + "\\TestData.xlsx", "DrugPricing", keep_default_na=False)
    df4 = pd.read_excel(c + "\\TestData.xlsx", "RTM", keep_default_na=False)
    df5 = pd.read_excel(c + "\\TestData.xlsx", "MedicareUDLSCoverage", dtype={"IN_ZipCode":"string"}, keep_default_na=False)
    df6 = pd.read_excel(c + "\\TestData.xlsx", "MedicareZipPlanValidation", dtype={"IN_ZipCode":"string"}, keep_default_na=False)



    df1.set_index("TestCaseId", drop=True, inplace=True)
    df2.set_index("TestCaseId", drop=True, inplace=True)
    df3.set_index("TestCaseId", drop=True, inplace=True)
    df4.set_index("TestCaseId", drop=True, inplace=True)
    df5.set_index("TestCaseId", drop=True, inplace=True)
    df6.set_index("TestCaseId", drop=True, inplace=True)



    test1 = df1.to_dict(orient="index")
    test2 = df2.to_dict(orient="index")
    test3 = df3.to_dict(orient="index")
    test4 = df4.to_dict(orient="index")
    test5 = df5.to_dict(orient="index")
    test6 = df6.to_dict(orient="index")


    f = open(c + "\\TestData.robot", "w")
    print("Saved to file")
    f.write("*** Variables ***\n")
    f.write("\n######## TS_Security ###################\n\n")
    f.write(str(test1))
    f.write("\n  ######## Smoke ################### \n\n")
    f.write(str(test2))
    f.write("\n  ######## DrugPricing ################### \n\n")
    f.write(str(test3))
    f.write("\n  ######## RTM ################### \n\n")
    f.write(str(test4))
    f.write("\n  ######## Medicare UDLS Coverage ################### \n\n")
    f.write(str(test5))
    f.write("\n  ######## Medicare UDLS ZipPlan Validation ################### \n\n")
    f.write(str(test6))
    f.close()


def format_data():
    fin = open("C:\\DCOE_RobotAutomationFramework_Azure\\Data\\TestData.robot", "rt")
    data = fin.read()
    data = re.sub("},", "\n", data)
    data = re.sub("}}", "\n", data)
    data = re.sub("{'", "", data)
    data = re.sub("':", "=", data)
    data = re.sub("}=", "} ", data)
    data = re.sub("= '", "=", data)
    data = re.sub("', '", "   ", data)
    data = re.sub(" '&", "&", data)
    data = re.sub("'\n", "\n", data)
    data = re.sub(", '", "  ", data)
    data = re.sub("= ", "=", data)
    data = re.sub('"', "", data)
    fin.close()
    fin = open("C:\\DCOE_RobotAutomationFramework_Azure\\Data\\TestData.robot", "wt")
    fin.write(data)
    fin.close()



def format_excel():
    workbook = xlsxwriter.Workbook('C:\\DCOE_RobotAutomationFramework_Azure\\Data\\output.xlsx')
    worksheet = workbook.add_worksheet('Sheet1')
    worksheet.add_table('A1:F5',{'style' : 'Table Style Light 17'})

def add_Application_Detail_into_Excel_TestData(Totalcount,Sheetname,ColNameArray,row_num,Parameter):
    wb = openpyxl.load_workbook('C:\\DCOE_RobotAutomationFramework_Azure\\Data\\TestData.xlsx')
    sheet = wb.get_sheet_by_name(Sheetname)
    for x in range(Totalcount):
        sheet[ColNameArray[x] + str(row_num)].value = Parameter[x]
        wb.save('C:\\DCOE_RobotAutomationFramework_Azure\\Data\\TestData.xlsx')